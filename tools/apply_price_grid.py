#!/usr/bin/env python3
"""
Load per-strength prices into the website.

From your edited workbook (Price Grid tab):
  python apply_price_grid.py DeliverMyMedications-Price-List.xlsx DeliverMyMedications-App.html

From a TelyRx crawl (telyrx_price_grid.py output), also refreshing the workbook:
  python apply_price_grid.py telyrx_price_grid.csv DeliverMyMedications-App.html [DeliverMyMedications-Price-List.xlsx]

Rules for crawled prices: GLP-1s = TelyRx minus $1; everything else = TelyRx price.
Compounded GLP-1s keep their slider and one flat price, so they're skipped.
"""
import csv, json, re, sys
GLP1 = re.compile(r"semaglutide|tirzepatide|liraglutide|exenatide|wegovy|ozempic|mounjaro|zepbound|rybelsus|foundayo", re.I)
GLP1_DISCOUNT = 1.00
slug = lambda s: re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def load_catalog(html):
    return json.loads(re.search(r"const CATALOG = (\[.*?\]);\n", html, re.S).group(1))

def add(grid, pid, dose, qty, price):
    g = grid.setdefault(pid, {"doses": [], "qtys": [], "prices": {}})
    if dose not in g["doses"]: g["doses"].append(dose)
    if qty not in g["qtys"]: g["qtys"].append(qty)
    g["prices"][f"{dose}|{qty}"] = round(float(price), 2)

def from_xlsx(path, cat):
    from openpyxl import load_workbook
    ids = {c["id"]: c for c in cat}
    ws = load_workbook(path, data_only=True)["Price Grid"]
    grid, bad = {}, []
    for row in ws.iter_rows(min_row=2, max_col=8, values_only=True):
        pid, _, name, dose, qty, tely, yours, used = row
        if not pid: continue
        c = ids.get(pid)
        if not c: bad.append(pid); continue
        if c.get("slider"): continue
        price = used if used is not None else (yours if yours is not None else tely)
        if price is None: continue
        add(grid, pid, str(dose), str(qty), price)
    # Products with a single row at the base price don't need a grid.
    for pid in list(grid):
        g = grid[pid]; c = ids[pid]
        if len(g["qtys"]) == 1 and g["doses"] == c["dosages"] and all(v == c["price"] for v in g["prices"].values()):
            del grid[pid]
    return grid, bad

def from_csv(path, cat):
    by_name = {slug(c["name"]): c for c in cat}
    grid, bad, rows = {}, set(), []
    for r in csv.DictReader(open(path)):
        c = by_name.get(slug(r["name"]))
        if not c or c.get("slider"): bad.add(r["name"]); continue
        price = float(r["price"])
        if GLP1.search(c["name"]): price -= GLP1_DISCOUNT
        d, q = r["strength"] or "Standard", r["quantity"] or "Standard quantity"
        add(grid, c["id"], d, q, price)
        rows.append((c["id"], c["cat"], c["name"], d, q, float(r["price"]), round(price, 2)))
    return grid, sorted(bad), rows

def write_workbook(path, rows):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path); ws = wb["Price Grid"]
    ws.delete_rows(2, ws.max_row)
    for i, row in enumerate(sorted(rows, key=lambda x: (x[1], x[2])), start=2):
        pid, catg, name, d, q, tely, yours = row
        ws.append([pid, catg, name, d, q, tely, yours, f'=IF(G{i}<>"",G{i},F{i})'])
        for c in ws[i]: c.font = Font(name="Arial")
        for col in "FGH": ws[f"{col}{i}"].number_format = '$#,##0.00'
    wb.save(path); print("Workbook Price Grid replaced with crawled prices. Open and save it in Excel to refresh totals.")

def main(src, html_path, xlsx=None):
    html = open(html_path).read()
    cat = load_catalog(html)
    if src.lower().endswith(".xlsx"):
        grid, bad = from_xlsx(src, cat)
    else:
        grid, bad, rows = from_csv(src, cat)
        if xlsx: write_workbook(xlsx, rows)
    html = re.sub(r"/\*GRID_START\*/.*?/\*GRID_END\*/",
        lambda m: "/*GRID_START*/const PRICE_GRID = " + json.dumps(grid, separators=(",", ":")) + ";/*GRID_END*/", html, flags=re.S)
    open(html_path, "w").write(html)
    print(f"Loaded per-strength prices for {len(grid)} products ({sum(len(g['prices']) for g in grid.values())} prices).")
    if bad: print(f"Skipped {len(bad)} unmatched rows: {', '.join(list(bad)[:12])}")

if __name__ == "__main__":
    if len(sys.argv) < 3: sys.exit(__doc__)
    main(*sys.argv[1:4])
